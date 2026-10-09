"""Typed spreadsheet facts and bounded retrieval projections.

Rows keep original coordinates. Merge inheritance is presentation only, never a
second fact. Formula evaluation/macros/external refresh are deliberately absent.
"""
from __future__ import annotations

import datetime
import hashlib
import json
import math
import re
import zipfile
import posixpath
from itertools import zip_longest
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from src import artifact_store, artifact_cache, table_batches
    from src.env_config import env_int
except ImportError:
    import artifact_store
    import artifact_cache
    import table_batches
    from env_config import env_int

INLINE_ROWS = env_int('PARSER_TABLE_INLINE_ROWS', 200)
INLINE_BYTES = env_int('PARSER_TABLE_INLINE_BYTES', 256 * 1024)
MAX_CELLS = env_int('PARSER_TABLE_MAX_CELLS', 5_000_000)
MAX_ROWS = env_int('PARSER_TABLE_MAX_ROWS', 1_000_000)
MAX_COLUMNS = env_int('PARSER_TABLE_MAX_COLUMNS', 1024)
MAX_XML_BYTES = env_int('PARSER_OFFICE_MAX_EXPANDED_BYTES', 512 * 1024 * 1024)
MAX_TABLES = env_int('PARSER_TABLE_MAX_REGIONS', 5000)
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def check_package(path: Path):
    if path.suffix.lower() not in {'.xlsx', '.docx', '.pptx'}:
        return
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 50_000 or sum(i.file_size for i in infos) > MAX_XML_BYTES:
            raise ValueError('Office package exceeds expanded input budget')
        if any(i.flag_bits & 1 for i in infos):
            raise ValueError('Encrypted Office packages are unsupported')


def chars(text) -> int:
    return len(re.sub(r'\s', '', str(text or '')))


def _value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, datetime.timedelta):
        return value.total_seconds()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _display(value, number_format: str) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'false' if not value else 'true'
    if isinstance(value, (datetime.date, datetime.time, datetime.timedelta)):
        return str(_value(value))
    # Preserve leading-zero identifiers and percentages; retain the format for
    # formats whose locale-dependent display cannot be reconstructed exactly.
    positive = number_format.split(';')[0]
    if isinstance(value, (int, float)) and re.fullmatch(r'0+', positive):
        return str(int(value)).zfill(len(positive)) if float(value).is_integer() else str(value)
    if isinstance(value, (int, float)) and re.fullmatch(r'0(?:\.0+)?%', positive):
        decimals = len(positive.split('.')[1].rstrip('%')) if '.' in positive else 0
        return f'{value * 100:.{decimals}f}%'
    if isinstance(value, float) and value.is_integer() and number_format == 'General':
        return str(int(value))
    return str(_value(value))


def _cell(cell, cached_cell):
    value, cached = cell.value, cached_cell.value
    formula = cell.data_type == 'f'
    effective = cached if formula else value
    if formula:
        kind = 'formula'
    elif effective is None:
        kind = 'blank'
    elif cell.data_type == 'e':
        kind = 'error'
    elif isinstance(effective, bool):
        kind = 'boolean'
    elif isinstance(effective, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        kind = 'date'
    elif isinstance(effective, (int, float)):
        kind = 'number'
    else:
        kind = 'string'
    result = {'coordinate': cell.coordinate, 'column': cell.column, 'type': kind,
              'value': _value(effective), 'display': _display(effective, cell.number_format),
              'number_format': cell.number_format}
    fmt = cell.number_format
    if '%' in fmt:
        result['unit'] = '%'
        result['display_scale'] = 100
    elif isinstance(effective, str) and re.fullmatch(r'[+-]?\d+(?:\.\d+)?%', effective.strip()):
        result['unit'] = '%'
        result['display_scale'] = 1
    else:
        currency = re.search(r'\[\$([^]\-]+)(?:-[^]]+)?\]|([$€£¥￥])', fmt)
        if currency:
            result['unit'] = currency.group(1) or currency.group(2)
            result['display_scale'] = 1
    if formula:
        result.update(formula=str(value), cached=_value(cached), cached_available=cached is not None)
        if cached is None:
            result['display'] = str(value)
    return result


def _escape(value):
    return str(value).replace('|', '\\|').replace('\r', ' ').replace('\n', '<br>')


class Region:
    def __init__(self, sheet, index, instance, hidden, epoch, merged_ranges, requested=None, definition=None):
        self.definition = definition
        self.id = f'sheet:{sheet}:table:{definition["name"]}' if definition else f'sheet:{sheet}:table:{index}'
        self.selected = not requested or f'sheet:{sheet}' in requested or self.id in requested
        self.header_rows = []
        self.header_hierarchy = {}
        self.previous_header_merged = False
        self.column_units = {}
        self.columns_seen = set()
        self.sheet, self.instance = sheet, instance
        self.rows = []
        self.preview_bytes = 0
        self.count = self.native_chars = 0
        self.minrow = self.maxrow = self.mincol = self.maxcol = None
        self.writer = None
        self.hidden, self.epoch, self.merged_ranges = hidden, epoch, merged_ranges
        self.headers = []
        self.header_columns = []

    def add(self, row):
        cells = row['cells']
        columns = [c['column'] for c in cells]
        self.columns_seen.update(columns)
        substantive = [c for c in cells if c['value'] is not None and not c.get('inherited')]
        if not self.count:
            substantive = [c for c in cells if c['value'] is not None and not c.get('inherited')]
            row['is_header'] = bool(self.definition and self.definition['header_count']) or (len(substantive) > 1 and all(c['type'] == 'string' for c in substantive))
            if not self.definition and all(c['type'] == 'string' for c in substantive) and any(c.get('merged_range') for c in cells):
                row['is_header'] = True
            self.header_columns = columns
            self.headers = [c['display'] if row['is_header'] else re.sub(r'\d', '', c['coordinate']) for c in cells]
        else:
            row['is_header'] = bool(not self.definition and self.previous_header_merged and substantive and all(c['type'] == 'string' for c in substantive))
        if row['is_header']:
            self.header_rows.append(row['row'])
            for cell in cells:
                if cell['display']:
                    chain = self.header_hierarchy.setdefault(cell['column'], [])
                    if not chain or chain[-1] != cell['display']:
                        chain.append(cell['display'])
            self.previous_header_merged = any(c.get('merged_range') for c in cells)
        else:
            self.previous_header_merged = False
        row['is_summary'] = bool(self.definition and self.definition['totals_count'] and row['row'] > self.definition['bounds'][3] - self.definition['totals_count'])
        # A vertical aggregate over preceding source rows is an explicit
        # calculation, unlike an unknown text label such as a business total.
        if not row['is_summary']:
            for cell in cells:
                match = re.fullmatch(r'=(?:SUM|AVERAGE|MIN|MAX|COUNT|COUNTA)\(\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)\)', cell.get('formula', ''), flags=re.I)
                if match and match[1] == match[3] and int(match[2]) <= int(match[4]) < row['row']:
                    row['is_summary'] = True
        for cell in cells:
            chain = self.header_hierarchy.get(cell['column'], [])
            declared = re.search(r'(?:\(([^()]+)\)|（([^（）]+)）)\s*$', chain[-1]) if chain else None
            if declared:
                cell['header_unit'] = declared.group(1) or declared.group(2)
                if not row['is_header'] and cell['type'] in {'number', 'formula'} and 'unit' not in cell:
                    cell.update(unit=cell['header_unit'], unit_source='header', display_scale=1)
            if not row['is_header'] and cell.get('unit'):
                self.column_units.setdefault(cell['column'], set()).add(cell['unit'])
        row.setdefault('hidden', False)
        self.count += 1
        self.native_chars += sum(chars(c['display']) for c in cells if not c.get('inherited'))
        self.minrow = min(self.minrow or row['row'], row['row'])
        self.maxrow = row['row']
        self.mincol = min(self.mincol or min(columns), min(columns))
        self.maxcol = max(self.maxcol or max(columns), max(columns))
        if not self.selected:
            return
        size = len(json.dumps(row, ensure_ascii=False).encode())
        if self.writer is None and (self.count > INLINE_ROWS or self.preview_bytes + size > INLINE_BYTES):
            self.writer = table_batches.BatchedWriter(self.instance, self.id)
            for previous in self.rows:
                self.writer.write(previous)
        if self.writer:
            self.writer.write(row)
        if self.count <= INLINE_ROWS and self.preview_bytes + size <= INLINE_BYTES:
            self.rows.append(row)
            self.preview_bytes += size

    def finish(self):
        from openpyxl.utils import get_column_letter
        if not self.count:
            return None, ''
        table_range = f'{get_column_letter(self.mincol)}{self.minrow}:{get_column_letter(self.maxcol)}{self.maxrow}'
        self.header_columns = sorted(self.columns_seen)
        self.headers = [' / '.join(self.header_hierarchy[c]) if c in self.header_hierarchy else get_column_letter(c) for c in self.header_columns]
        header_units = []
        for column in self.header_columns:
            chain = self.header_hierarchy.get(column, [])
            declared = re.search(r'(?:\(([^()]+)\)|（([^（）]+)）)\s*$', chain[-1]) if chain else None
            header_units.append((declared.group(1) or declared.group(2)) if declared else None)
        table = {'id': self.id, 'sheet': self.sheet, 'range': table_range, 'headers': self.headers,
                 'header_columns': self.header_columns, 'header_units': header_units, 'column_units': {str(c): sorted(units) for c, units in self.column_units.items()}, 'header_rows': self.header_rows, 'header_hierarchy': self.header_hierarchy, 'rows': self.rows, 'row_count': self.count,
                 'complete': True, 'mode': 'artifact' if self.writer else 'inline',
                 'hidden': self.hidden, 'date_system': self.epoch,
                 'merged_ranges': self.merged_ranges, 'formula_capability': 'formula-and-cache',
                 'layout': 'defined-table' if self.definition else 'inferred-region',
                 'warnings': [] if self.definition else ['Table boundaries/header roles are inferred from source structure; validate unspecified summary rows and units before whole-table aggregation'],
                 'summary_detection': 'defined-totals-or-vertical-aggregate-formula',
                 'range_origin': {'row': self.minrow, 'column': self.mincol}}
        if self.writer:
            table['artifact_id'] = self.writer.finish()
            table['stream_batches'] = len(self.writer.batches)
            table['reused_batches'] = self.writer.reused
            table['batch_rows'] = table_batches.BATCH_ROWS
        width = self.maxcol - self.mincol + 1
        if width > MAX_COLUMNS:
            raise ValueError('Table exceeds column budget')
        # First source row is retained even when it is not a header.
        header = self.headers if len(self.headers) == width else [get_column_letter(c) for c in range(self.mincol, self.maxcol + 1)]
        lines = [f'### 工作表：{self.sheet} · {table_range}',
                 '| ' + ' | '.join(_escape(h) for h in header) + ' |',
                 '| ' + ' | '.join(['---'] * width) + ' |']
        for row in self.rows:
            if row['is_header']:
                continue
            mapped = {c['column']: c['display'] for c in row['cells']}
            lines.append('| ' + ' | '.join(_escape(mapped.get(c, '')) for c in range(self.mincol, self.maxcol + 1)) + ' |')
        if self.writer:
            lines.append(f'<!-- structured-table:{self.id}; rows:{self.count}; preview:{len(self.rows)} -->')
        return table, '\n'.join(lines)


def _xlsx_merges(path, sheet_paths):
    from openpyxl.utils.cell import range_boundaries
    results = {}
    with zipfile.ZipFile(path) as archive:
        for name, xmlpath in sheet_paths.items():
            ranges, hidden_rows, hidden_columns, references = [], set(), [], []
            xmlpath = xmlpath.lstrip('/')
            with archive.open(xmlpath) as source:
                for event, node in ET.iterparse(source, events=('end',)):
                    if node.tag == NS + 'mergeCell':
                        ranges.append((node.attrib['ref'], range_boundaries(node.attrib['ref'])))
                        if len(ranges) > 100_000:
                            raise ValueError('Workbook exceeds merge complexity budget')
                    elif node.tag == NS + 'row' and node.attrib.get('hidden') == '1':
                        hidden_rows.add(int(node.attrib['r']))
                    elif node.tag == NS + 'tablePart':
                        references.append(node.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
                    elif node.tag == NS + 'col' and node.attrib.get('hidden') == '1':
                        hidden_columns.append((int(node.attrib['min']), int(node.attrib['max'])))
                    node.clear()
            definitions = []
            relation_path = posixpath.join(posixpath.dirname(xmlpath), '_rels', posixpath.basename(xmlpath) + '.rels')
            if references and relation_path in archive.namelist():
                relationships = ET.fromstring(archive.read(relation_path))
                for relation in relationships:
                    if relation.attrib.get('Id') not in references or relation.attrib.get('TargetMode') == 'External':
                        continue
                    target = relation.attrib['Target']
                    target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(xmlpath), target))
                    if not target.startswith('xl/tables/') or target not in archive.namelist():
                        raise ValueError('Invalid workbook table relationship')
                    node = ET.fromstring(archive.read(target))
                    definitions.append({'name': node.attrib.get('displayName') or node.attrib['name'],
                        'bounds': range_boundaries(node.attrib['ref']), 'header_count': int(node.attrib.get('headerRowCount', 1)),
                        'totals_count': int(node.attrib.get('totalsRowCount', 0))})
            results[name] = (ranges, hidden_rows, hidden_columns, definitions)
    return results


def extract(path: Path, instance='legacy', unit_ids=None):
    check_package(path)
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    token = table_batches.source_hash.set(digest.hexdigest())
    identity_token = artifact_cache.instance_identity.set(instance)
    try:
        return _xls(path, instance, unit_ids) if path.suffix.lower() == '.xls' else _xlsx(path, instance, unit_ids)
    finally:
        table_batches.source_hash.reset(token)
        artifact_cache.instance_identity.reset(identity_token)


def _result(tables, units, sections, native_chars):
    return {'structured_tables': tables, 'source_units': units, 'markdown': '\n\n'.join(sections),
            'native_text_chars': native_chars, 'generated_text_chars': 0,
            'coverage': {'total': len(units), 'processed': sum(u['status'] == 'processed' for u in units),
                         'failed': sum(u['status'] == 'failed' for u in units),
                         'skipped': sum(u['status'] == 'skipped' for u in units)}}


def _xlsx(path, instance, unit_ids):
    import openpyxl
    from openpyxl.utils.cell import get_column_letter
    # Two read-only iterators preserve formula and cached result independently.
    # Package expansion and process memory are bounded before loading strings.
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    active = []
    try:
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        paths = {sheet.title: sheet._worksheet_path for sheet in formulas.worksheets}
        metadata = _xlsx_merges(path, paths)
        tables, units, sections = [], [], []
        native_chars = cell_count = 0
        for sheet in formulas.worksheets:
            sheetid = f'sheet:{sheet.title}'
            selected_sheet = not unit_ids or sheetid in unit_ids or any(u.startswith(sheetid + ':') for u in unit_ids)
            if not selected_sheet:
                units.append({'id': sheetid, 'kind': 'sheet', 'sheet': sheet.title, 'status': 'skipped',
                              'native_text_chars': 0, 'generated_text_chars': 0, 'error': 'not_selected'})
                continue
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError('Worksheet exceeds row/column budget')
            merges, hidden_rows, hidden_columns, definitions = metadata[sheet.title]
            defined_regions = {}
            anchors = {}
            region = None
            index = 0
            # Separate regions on empty rows; original row indices never shift.
            for number, pair in enumerate(zip_longest(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                raw, cache = pair
                if number > MAX_ROWS:
                    raise ValueError('Worksheet exceeds row budget')
                cells = []
                for col, (cell, cached_cell) in enumerate(zip_longest(raw or [], cache or []), 1):
                    if cell is None or cell.value is None:
                        continue
                    item = _cell(cell, cached_cell or cell)
                    if number in hidden_rows or any(lo <= col <= hi for lo, hi in hidden_columns):
                        item['hidden'] = True
                    anchors[item['coordinate']] = item if any(bounds[0] == col and bounds[1] == number for _, bounds in merges) else anchors.get(item['coordinate'])
                    if anchors.get(item['coordinate']) is None:
                        anchors.pop(item['coordinate'], None)
                    cells.append(item)
                mapped = {c['column']: c for c in cells}
                for label, (lo_col, lo_row, hi_col, hi_row) in merges:
                    if not lo_row <= number <= hi_row:
                        continue
                    coordinate = f'{get_column_letter(lo_col)}{lo_row}'
                    anchor = anchors.get(coordinate)
                    if not anchor:
                        continue
                    for col in range(lo_col, hi_col + 1):
                        if col == lo_col and number == lo_row:
                            anchor['merged_anchor'] = coordinate
                            anchor['merged_range'] = label
                        elif col not in mapped:
                            mapped[col] = {'coordinate': f'{get_column_letter(col)}{number}', 'column': col,
                                           'type': 'blank', 'value': None, 'display': anchor['display'],
                                           'inherited': True, 'merged_anchor': coordinate, 'merged_range': label}
                cells = sorted(mapped.values(), key=lambda c: c['column'])
                cell_count += len(cells)
                if cell_count > MAX_CELLS:
                    raise ValueError('Workbook exceeds effective cell budget')
                # Explicit Excel tables can sit side by side. Route cells by
                # their original rectangles; keep non-table notes in separate
                # inferred regions instead of discarding them.
                remaining = {cell['column']: cell for cell in cells}
                for definition in definitions:
                    lo_col, lo_row, hi_col, hi_row = definition['bounds']
                    if not lo_row <= number <= hi_row:
                        continue
                    subset = [cell for cell in cells if lo_col <= cell['column'] <= hi_col]
                    if not subset:
                        continue
                    identifier = definition['name']
                    target = defined_regions.get(identifier)
                    if target is None:
                        target = Region(sheet.title, identifier, instance, sheet.sheet_state != 'visible',
                            '1904' if formulas.epoch.year == 1904 else '1900', [m[0] for m in merges], unit_ids, definition)
                        defined_regions[identifier] = target
                        active.append(target)
                    target.add({'row': number, 'cells': subset, 'hidden': number in hidden_rows})
                    for cell in subset:
                        remaining.pop(cell['column'], None)
                cells = list(remaining.values())
                if not cells:
                    if region:
                        table, md = region.finish()
                        _append_region(region, table, md, unit_ids, tables, units, sections)
                        native_chars += region.native_chars if not unit_ids or sheetid in unit_ids or region.id in unit_ids else 0
                        active.remove(region)
                        region = None
                    continue
                if region is None:
                    index += 1
                    if len(tables) + index > MAX_TABLES:
                        raise ValueError('Workbook exceeds table region budget')
                    region = Region(sheet.title, index, instance, sheet.sheet_state != 'visible',
                                    '1904' if formulas.epoch.year == 1904 else '1900', [m[0] for m in merges], unit_ids)
                    active.append(region)
                region.add({'row': number, 'cells': cells, 'hidden': number in hidden_rows})
            if region:
                table, md = region.finish()
                _append_region(region, table, md, unit_ids, tables, units, sections)
                native_chars += region.native_chars if not unit_ids or sheetid in unit_ids or region.id in unit_ids else 0
                active.remove(region)
            for target in defined_regions.values():
                table, md = target.finish()
                _append_region(target, table, md, unit_ids, tables, units, sections)
                native_chars += target.native_chars if target.selected else 0
                active.remove(target)
            if not index and not defined_regions:
                units.append({'id': sheetid, 'kind': 'sheet', 'sheet': sheet.title, 'status': 'skipped',
                              'native_text_chars': 0, 'generated_text_chars': 0, 'error': 'empty_sheet'})
        return _result(tables, units, sections, native_chars)
    finally:
        for region in active:
            if region.writer:
                region.writer.abort()
        formulas.close()
        if cached:
            cached.close()


def _append_region(region, table, md, requested, tables, units, sections):
    selected = not requested or f'sheet:{region.sheet}' in requested or region.id in requested
    units.append({'id': region.id, 'kind': 'table', 'sheet': region.sheet,
                  'range': table['range'], 'table_id': table['id'], 'status': 'processed' if selected else 'skipped',
                  'native_text_chars': region.native_chars if selected else 0, 'generated_text_chars': 0, 'source_kind': 'native', 'markdown': md if selected else ''})
    if selected:
        tables.append(table)
        sections.append(md)
    elif table.get('artifact_id'):
        # Unselected region can be traversed to determine stable region IDs,
        # but no artifact should linger or be downloaded for it.
        artifact_store.remove(table['artifact_id'], region.instance)


def _xls(path, instance, unit_ids):
    import xlrd
    from openpyxl.utils.cell import get_column_letter
    workbook = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    tables, units, sections = [], [], []
    native_chars = count = 0
    active = None
    try:
        for sheet in workbook.sheets():
            sheetid = f'sheet:{sheet.name}'
            if unit_ids and sheetid not in unit_ids and not any(u.startswith(sheetid + ':') for u in unit_ids):
                units.append({'id': sheetid, 'kind': 'sheet', 'sheet': sheet.name, 'status': 'skipped',
                              'native_text_chars': 0, 'generated_text_chars': 0})
                continue
            if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                raise ValueError('Worksheet exceeds row/column budget')
            index = 0
            merge_labels = [f'{get_column_letter(clo + 1)}{rlo + 1}:{get_column_letter(chi)}{rhi}' for rlo, rhi, clo, chi in sheet.merged_cells]
            for r in range(sheet.nrows):
                cells = []
                for c in range(sheet.ncols):
                    raw = sheet.cell(r, c)
                    if raw.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        continue
                    value = raw.value
                    kind = {xlrd.XL_CELL_TEXT: 'string', xlrd.XL_CELL_NUMBER: 'number',
                            xlrd.XL_CELL_DATE: 'date', xlrd.XL_CELL_BOOLEAN: 'boolean', xlrd.XL_CELL_ERROR: 'error'}.get(raw.ctype, 'string')
                    if kind == 'boolean':
                        value = bool(value)
                    elif kind == 'date':
                        value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                    elif kind == 'error':
                        value = xlrd.error_text_from_code.get(value, str(value))
                    fmt = 'General'
                    if raw.xf_index is not None:
                        fmt = workbook.format_map[workbook.xf_list[raw.xf_index].format_key].format_str
                    cells.append({'coordinate': f'{get_column_letter(c + 1)}{r + 1}', 'column': c + 1,
                                  'type': kind, 'value': _value(value), 'display': _display(value, fmt), 'number_format': fmt})
                mapped = {cell['column']: cell for cell in cells}
                for rlo, rhi, clo, chi in sheet.merged_cells:
                    if not rlo <= r < rhi:
                        continue
                    anchor_value = sheet.cell_value(rlo, clo)
                    anchor_coordinate = f'{get_column_letter(clo + 1)}{rlo + 1}'
                    for c in range(clo, chi):
                        if r == rlo and c == clo and c + 1 in mapped:
                            mapped[c + 1]['merged_anchor'] = anchor_coordinate
                        elif c + 1 not in mapped:
                            mapped[c + 1] = {'coordinate': f'{get_column_letter(c + 1)}{r + 1}', 'column': c + 1,
                                             'type': 'blank', 'value': None, 'display': str(anchor_value) if anchor_value is not None else '',
                                             'inherited': True, 'merged_anchor': anchor_coordinate}
                cells = sorted(mapped.values(), key=lambda cell: cell['column'])
                count += len(cells)
                if count > MAX_CELLS:
                    raise ValueError('Workbook exceeds effective cell budget')
                if not cells:
                    if active:
                        table, md = active.finish()
                        table['formula_capability'] = 'cached-only'
                        _append_region(active, table, md, unit_ids, tables, units, sections)
                        native_chars += active.native_chars if not unit_ids or sheetid in unit_ids or active.id in unit_ids else 0
                        active = None
                    continue
                if not active:
                    index += 1
                    active = Region(sheet.name, index, instance, getattr(sheet, 'visibility', 0) != 0,
                                    '1904' if workbook.datemode else '1900', merge_labels, unit_ids)
                active.add({'row': r + 1, 'cells': cells})
            if active:
                table, md = active.finish()
                table['formula_capability'] = 'cached-only'
                _append_region(active, table, md, unit_ids, tables, units, sections)
                native_chars += active.native_chars if not unit_ids or sheetid in unit_ids or active.id in unit_ids else 0
                active = None
            workbook.unload_sheet(sheet.name)
        return _result(tables, units, sections, native_chars)
    finally:
        if active and active.writer:
            active.writer.abort()
        workbook.release_resources()
