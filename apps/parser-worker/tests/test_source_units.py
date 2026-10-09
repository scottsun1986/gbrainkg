"""Source placement, real-content policy, image coverage and retry fixtures."""
import asyncio
import io
import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import docx
import httpx
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

import main
import image_units
import source_contract
import artifact_store

artifact_store = main.artifact_store


def png(size=(300, 120), color='white'):
    buffer = io.BytesIO()
    image = Image.new('RGB', size, color)
    image.save(buffer, format='PNG')
    return buffer.getvalue()


class SourceUnitTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for module, name, value in [(main, 'UPLOAD_ROOT', self.root),
                                    (artifact_store, 'ROOT', self.root / 'assets')]:
            patcher = patch.object(module, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_docx_inline_image_anchor_and_stable_image_retry_ids(self):
        document = docx.Document()
        document.add_paragraph('步骤一：打开菜单。')
        document.add_picture(io.BytesIO(png()), width=docx.shared.Inches(1))
        document.add_paragraph('步骤二：选择按钮。')
        document.add_picture(io.BytesIO(png(color='gray')), width=docx.shared.Inches(1))
        document.sections[0].header.paragraphs[0].text = '文档版本说明'
        path = self.root / 'steps.docx'
        document.save(path)
        full = {}
        text, images = main.extract_docx(path, full)
        self.assertLess(text.index('打开菜单'), text.index('docx-media-1'))
        self.assertLess(text.index('docx-media-1'), text.index('选择按钮'))
        self.assertEqual(images[1]['anchor'], 'docx:block:4')
        self.assertTrue(any(u['kind'] == 'header' and u['markdown'] == '文档版本说明' for u in full['source_units']))
        partial = {}
        projection, retry_images = main.extract_docx(path, partial, ['docx-media-2'])
        self.assertEqual([image['key'] for image in retry_images], ['docx-media-2'])
        self.assertNotIn('打开菜单', projection)
        self.assertTrue(all(u['status'] == 'skipped' for u in partial['source_units']))

    async def test_repeated_asset_hash_recognised_once_keeps_both_positions(self):
        blob = png()
        images = [{'key': 'a', 'blob': blob, 'ext': 'png', 'anchor': 'paragraph:1'},
                  {'key': 'b', 'blob': blob, 'ext': 'png', 'anchor': 'paragraph:2'}]
        recognizer = AsyncMock(return_value=('图片中的确认按钮', {'ocr_words_result_num': 1}))
        with patch.object(main, 'convert_image_with_baidu_ocr', recognizer):
            text, metadata = await main.ocr_image_parts_into_markdown('<!-- image: a -->\n\n正文\n\n<!-- image: b -->', images, {'provider': 'baidu'})
        self.assertEqual(recognizer.await_count, 1)
        self.assertEqual(text.count('图片中的确认按钮'), 2)
        self.assertEqual(metadata['assets'][0]['sha256'], metadata['assets'][1]['sha256'])
        self.assertEqual([u['anchor'] for u in metadata['source_units']], ['paragraph:1', 'paragraph:2'])
        self.assertEqual(metadata['image_hash_reuses'], 1)
        self.assertTrue(all(u['source_kind'] == 'ocr' for u in metadata['source_units']))

    async def test_vlm_is_derived_content_and_failed_images_are_visible(self):
        with patch.object(main, 'is_vlm_available', return_value=True), patch.object(main, 'describe_image_with_vlm', AsyncMock(return_value='箭头连接左右两个节点')):
            text, metadata = await main.ocr_image_parts_into_markdown('<!-- image: flow -->', [{'key': 'flow', 'blob': png(), 'ext': 'png'}], {'provider': 'none'})
        self.assertIn('模型派生', text)
        self.assertEqual(metadata['native_text_chars'], 0)
        self.assertGreater(metadata['generated_text_chars'], 0)
        self.assertEqual(metadata['assets'][0]['source_kind'], 'visual')
        with patch.object(main, 'is_vlm_available', return_value=False):
            _, metadata = await main.ocr_image_parts_into_markdown('<!-- image: pending -->', [{'key': 'pending', 'blob': png(), 'ext': 'png'}], {'provider': 'none'})
        self.assertEqual(metadata['source_units'][0]['status'], 'failed')
        self.assertEqual(metadata['source_units'][0]['native_text_chars'], 0)

    def test_pptx_text_chart_notes_and_picture_positions(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(4), Inches(1)).text_frame.text = '来源幻灯片'
        slide.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(1.5), Inches(2), Inches(1))
        data = CategoryChartData()
        data.categories = ['A', 'B']
        data.add_series('原始数量', (0, 12))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(2), data)
        slide.notes_slide.notes_text_frame.text = '演讲者原始备注'
        path = self.root / 'source.pptx'
        presentation.save(path)
        contract = {}
        blocks, images = main.extract_pptx_native(path, contract)
        self.assertIn('原始数量', blocks[0])
        self.assertIn('0.0', blocks[0])
        self.assertIn('演讲者原始备注', blocks[0])
        self.assertLess(blocks[0].index('来源幻灯片'), blocks[0].index('<!-- image:'))
        self.assertTrue(images[0]['bbox'])
        self.assertTrue(any(u['kind'] == 'chart' for u in contract['source_units']))
        self.assertTrue(any(u['kind'] == 'notes' for u in contract['source_units']))

    @unittest.skipUnless(main.PYMUPDF_INSTALLED, 'PDF region test needs optional pdf-render extra')
    async def test_pdf_text_page_screenshot_scan_page_failure_and_local_retry(self):
        import fitz
        document = fitz.open()
        first = document.new_page()
        first.insert_text((50, 50), 'Original native instructions ' * 8)
        first.insert_image(fitz.Rect(50, 90, 350, 210), stream=png())
        second = document.new_page()
        second.insert_image(fitz.Rect(0, 0, 600, 750), stream=png((600, 750), 'gray'))
        path = self.root / 'mixed.pdf'
        document.save(path)
        document.close()
        info = main.inspect_pdf_native(path)
        with patch.object(main, 'convert_image_with_baidu_ocr', AsyncMock(return_value=('Screen label', {'ocr_words_result_num': 1}))), patch.object(main, 'convert_with_cloud_ocr', AsyncMock(side_effect=RuntimeError('Unavailable'))), patch.object(main, 'LOCAL_DOCLING_ENABLED', False):
            markdown, _, contract = await main.convert_pdf_structured(path, info, {'provider': 'baidu'})
        self.assertIn('Original native instructions', markdown)
        self.assertIn('Screen label', markdown)
        self.assertTrue(any(u['kind'] == 'image' and u['page'] == 1 and u.get('bbox') for u in contract['source_units']))
        # The screenshot recogniser may rescue a scan page after its document
        # OCR fails, but the page-level failure remains visible.
        self.assertTrue(any(u['id'] == 'page:2' and u['status'] == 'failed' for u in contract['source_units']))
        with patch.object(main, 'convert_image_with_baidu_ocr', AsyncMock(return_value=('Screen label', {}))), patch.object(main, 'LOCAL_DOCLING_ENABLED', False):
            partial, _, retry = await main.convert_pdf_structured(path, info, {'provider': 'baidu'}, ['page:1'])
        self.assertIn('第 1 页', partial)
        self.assertNotIn('第 2 页', partial)
        self.assertTrue(any(u['id'] == 'page:2' and u['status'] == 'skipped' for u in retry['source_units']))
        main.cleanup_native(path)

    def test_png_webp_bmp_tiff_frames_pixels_and_long_image_tiles(self):
        for extension, format_name in [('png', 'PNG'), ('webp', 'WEBP'), ('bmp', 'BMP')]:
            path = self.root / ('source.' + extension)
            Image.new('RGB', (120, 80), 'white').save(path, format=format_name)
            parts, skipped = image_units.prepare(path)
            self.assertEqual(len(parts), 1)
            self.assertEqual(parts[0]['page'], 1)
            self.assertFalse(skipped)
        path = self.root / 'multi.tiff'
        Image.new('RGB', (120, 80), 'white').save(path, save_all=True, append_images=[Image.new('RGB', (120, 80), 'gray')])
        parts, _ = image_units.prepare(path)
        self.assertEqual([p['page'] for p in parts], [1, 2])
        retry, skipped = image_units.prepare(path, ['image:frame:2'])
        self.assertEqual([p['page'] for p in retry], [2])
        self.assertEqual(skipped[0]['status'], 'skipped')
        with patch.object(image_units, 'MAX_PIXELS', 100):
            with self.assertRaisesRegex(ValueError, 'pixel budget'):
                image_units.prepare(path)
        long = self.root / 'long.png'
        Image.new('RGB', (300, 6000), 'white').save(long)
        parts, _ = image_units.prepare(long)
        self.assertGreater(len(parts), 1)
        self.assertGreater(parts[0]['bbox'][1] + parts[0]['bbox'][3], parts[1]['bbox'][1])
        self.assertTrue(all(part['overlap_pixels'] > 0 for part in parts))

    async def test_execute_rejects_presentation_scaffolding_and_original_hash_mismatch(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(1))
        buffer = io.BytesIO()
        prs.save(buffer)
        with patch.dict(os.environ, {'AUTH_TOKEN': 'test-token'}), patch.object(main, 'LOCAL_DOCLING_ENABLED', False), patch.object(main, 'is_vlm_available', return_value=False):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://worker') as client:
                result = await client.post('/parse-execute', headers={'Authorization': 'Bearer test-token'}, files={'file': ('only.pptx', buffer.getvalue())}, data={'instance_id': 'inst1', 'ocr_provider': 'none'})
                self.assertEqual(result.json()['status'], 'failed')
                self.assertEqual(result.json()['content_text_chars'], 0)
                self.assertGreater(result.json()['coverage']['failed'], 0)
                mismatch = await client.post('/parse-execute', headers={'Authorization': 'Bearer test-token'}, files={'file': ('source.txt', b'original body')}, data={'instance_id': 'inst1', 'cache_source_hash': '0' * 64})
                self.assertEqual(mismatch.status_code, 409)

    async def test_legacy_ppt_guard_and_missing_converter_are_explicit(self):
        path = self.root / 'not-a-presentation.ppt'
        path.write_bytes(b'PK renamed office file')
        with self.assertRaisesRegex(ValueError, 'OLE2'):
            await main.convert_legacy_ppt(path, {}, {})
        path.write_bytes(bytes.fromhex('D0CF11E0A1B11AE1'))
        with patch.object(main.shutil, 'which', return_value=None):
            with self.assertRaisesRegex(RuntimeError, 'soffice'):
                await main.convert_legacy_ppt(path, {}, {})

    @unittest.skipUnless(shutil.which('soffice'), 'Real legacy-PPT conversion needs local soffice')
    async def test_real_legacy_ppt_cold_conversion_preserves_body(self):
        source = self.root / 'original.pptx'
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = 'Legacy slide source evidence'
        prs.save(source)
        profile = self.root / 'fixture-office-profile'
        subprocess.run(['soffice', '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'ppt', '--outdir', str(self.root), str(source)], check=True, timeout=120, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        legacy = self.root / 'original.ppt'
        with patch.object(main, 'is_vlm_available', return_value=False):
            text, _, metadata = await main.convert_legacy_ppt(legacy, {'provider': 'none'}, {'filename': 'original.ppt'})
        self.assertIn('Legacy slide source evidence', text)
        self.assertEqual(metadata['conversion_policy']['macros'], 'disabled')

    def test_plaintext_utf16_and_markdown_original_structure(self):
        text = '# 标题\n\n- 列表\n\n```python\nprint(0)\n```\n\n[^1]: 注释'
        self.assertEqual(main.extract_plaintext('source.md', text.encode('utf-16')), text)
        self.assertEqual(main.extract_plaintext('source.txt', '旧编码正文'.encode('gb18030')), '旧编码正文')
        skeleton = '# upload-title\n\n## 第 1 页\n\n<!-- image: a -->\n*(图片未提取到正文，待重试)*'
        self.assertFalse(source_contract.body_text(skeleton))
