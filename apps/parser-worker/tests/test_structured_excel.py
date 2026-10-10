"""Real spreadsheet fixtures: typed facts, complete streams and recovery."""
import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet.table import Table
import artifact_cache
import artifact_store
import structured_excel
import table_batches

artifact_store = structured_excel.artifact_store
artifact_cache = structured_excel.artifact_cache
table_batches = structured_excel.table_batches


class StructuredExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        for module, name, value in [(artifact_store, 'ROOT', self.root / 'artifacts'),
                                    (artifact_cache, 'ROOT', self.root / 'cache')]:
            patcher = patch.object(module, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def save(self, workbook, name='source.xlsx'):
        path = self.root / name
        workbook.save(path)
        workbook.close()
        return path

    def save_legacy(self, workbook, name='legacy.xls'):
        path = self.root / name
        workbook.save(str(path))
        return path

    def rewrite_sheet_xml(self, path, transform):
        with zipfile.ZipFile(path) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        key = 'xl/worksheets/sheet1.xml'
        entries[key] = transform(entries[key])
        with zipfile.ZipFile(path, 'w') as output:
            for name, data in entries.items():
                output.writestr(name, data)

    def set_dimension(self, path, dimension):
        def replace(xml):
            updated, count = re.subn(rb'<dimension ref="[^"]+"',
                                    b'<dimension ref="' + dimension.encode() + b'"', xml)
            self.assertEqual(count, 1)
            return updated
        self.rewrite_sheet_xml(path, replace)

    def test_incorrect_dimensions_do_not_truncate_spreadsheet_content(self):
        # A workbook exporter can write a plausible but incorrect dimension.
        # The complete 100 question/answer rows must reach the text projection.
        for dimension in ['A1', 'A1:B2', 'A1:XFD1048576']:
            with self.subTest(dimension=dimension):
                book = openpyxl.Workbook()
                sheet = book.active
                sheet.title = '产品问答'
                sheet.append(['问题', '答案'])
                for index in range(1, 101):
                    sheet.append([f'产品{index}的功能是什么？', f'产品{index}可以管理家庭设备。'])
                path = self.save(book)
                self.set_dimension(path, dimension)
                result = structured_excel.extract(path, 'inst1')
                table = result['structured_tables'][0]
                self.assertEqual(table['row_count'], 101)
                self.assertEqual(table['range'], 'A1:B101')
                self.assertTrue(table['complete'])
                self.assertEqual(len(table['rows']), 101)
                self.assertEqual(sum(len(row['cells']) for row in table['rows']), 202)
                self.assertIn('产品73的功能是什么？', result['markdown'])
                self.assertIn('产品73可以管理家庭设备。', result['markdown'])
                self.assertIn('产品100可以管理家庭设备。', result['markdown'])
                self.assertEqual(result['source_units'][0]['range'], 'A1:B101')

    def test_incorrect_dimensions_keep_formula_and_cached_streams_aligned(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['名称', '数值'])
        sheet.append(['已缓存', '=1+1'])
        sheet.append(['未缓存', '=2+2'])
        path = self.save(book)
        self.set_dimension(path, 'A1')
        self.rewrite_sheet_xml(path, lambda xml: xml.replace(b'<f>1+1</f><v></v>', b'<f>1+1</f><v>2</v>'))
        table = structured_excel.extract(path, 'inst1')['structured_tables'][0]
        cells = {cell['coordinate']: cell for row in table['rows'] for cell in row['cells']}
        self.assertEqual(table['range'], 'A1:B3')
        self.assertEqual(cells['B2']['formula'], '=1+1')
        self.assertEqual(cells['B2']['cached'], 2)
        self.assertEqual(cells['B2']['value'], 2)
        self.assertTrue(cells['B2']['cached_available'])
        self.assertEqual(cells['B3']['formula'], '=2+2')
        self.assertIsNone(cells['B3']['value'])
        self.assertFalse(cells['B3']['cached_available'])
        self.assertEqual(cells['B3']['display'], '')

    def test_incorrect_dimensions_cannot_bypass_actual_coordinate_budgets(self):
        for coordinate, budget in [('A20', {'MAX_ROWS': 10}), ('T1', {'MAX_COLUMNS': 10})]:
            with self.subTest(coordinate=coordinate):
                book = openpyxl.Workbook()
                sheet = book.active
                sheet['A1'] = '表头'
                sheet[coordinate] = '超限数据'
                path = self.save(book)
                self.set_dimension(path, 'A1')
                with patch.multiple(structured_excel, **budget), \
                        patch('openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows') as rows:
                    with self.assertRaisesRegex(ValueError, 'row.*budget'):
                        structured_excel.extract(path, 'inst1')
                    rows.assert_not_called()

    def test_unselected_worksheet_is_not_subject_to_actual_coordinate_budgets(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '保留'
        sheet.append(['名称', '金额'])
        sheet.append(['甲', 1])
        other = book.create_sheet('跳过')
        other['T20'] = '未选择数据'
        with patch.multiple(structured_excel, MAX_ROWS=10, MAX_COLUMNS=10):
            result = structured_excel.extract(self.save(book), 'inst1', ['sheet:保留'])
        self.assertEqual(len(result['structured_tables']), 1)
        self.assertEqual(result['structured_tables'][0]['range'], 'A1:B2')
        self.assertEqual(result['coverage']['skipped'], 1)

    def test_incorrect_dimensions_cannot_bypass_merge_range_budgets(self):
        for region, budget in [('A1:T1', {'MAX_COLUMNS': 10}), ('A1:A20', {'MAX_ROWS': 10})]:
            with self.subTest(region=region):
                book = openpyxl.Workbook()
                sheet = book.active
                sheet['A1'] = '合并单元格'
                path = self.save(book)
                self.set_dimension(path, 'A1')
                # Keep source cells at A1; only the merge range is oversized.
                self.rewrite_sheet_xml(path, lambda xml: xml.replace(
                    b'</worksheet>', b'<mergeCells count="1"><mergeCell ref="' +
                    region.encode() + b'"/></mergeCells></worksheet>'))
                with patch.multiple(structured_excel, **budget), \
                        patch('openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows') as rows:
                    with self.assertRaisesRegex(ValueError, 'row/column budget'):
                        structured_excel.extract(path, 'inst1')
                    rows.assert_not_called()

    def test_incorrect_dimensions_cannot_bypass_defined_table_range_budgets(self):
        for region, budget in [('A1:T2', {'MAX_COLUMNS': 10}), ('A1:B20', {'MAX_ROWS': 10})]:
            with self.subTest(region=region):
                book = openpyxl.Workbook()
                sheet = book.active
                sheet.append(['名称', '数值'])
                sheet.append(['甲', 1])
                sheet.add_table(Table(displayName='Facts', ref='A1:B2'))
                path = self.save(book)
                self.set_dimension(path, 'A1')
                # Only the table definition exceeds the budget; source cells
                # remain in A1:B2, so this exercises the structure range guard.
                with zipfile.ZipFile(path) as source:
                    entries = {name: source.read(name) for name in source.namelist()}
                key = 'xl/tables/table1.xml'
                entries[key] = entries[key].replace(b'ref="A1:B2"', b'ref="' + region.encode() + b'"')
                with zipfile.ZipFile(path, 'w') as output:
                    for name, data in entries.items():
                        output.writestr(name, data)
                with patch.multiple(structured_excel, **budget), \
                        patch('openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows') as rows:
                    with self.assertRaisesRegex(ValueError, 'row/column budget'):
                        structured_excel.extract(path, 'inst1')
                    rows.assert_not_called()

    def test_zero_false_formula_cache_merge_coordinates_units_and_hidden(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '事实'
        sheet.append(['金额（万元）', '开关', '编号', '比例（%）', '公式', '父类'])
        sheet.append([0, False, '00123', .15, '=A2+1', '分组'])
        sheet.append([12, True, '00124', .25, '=A3+1', None])
        sheet['D2'].number_format = '0%'
        sheet['A3'].number_format = '[$¥-804]#,##0.00'
        sheet.merge_cells('F2:F3')
        sheet.row_dimensions[3].hidden = True
        path = self.save(book)
        # openpyxl writes no cached formula result; inject a real cached result
        # in just E3 and prove that E2 remains explicitly unavailable.
        with zipfile.ZipFile(path) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(b'<f>A3+1</f><v></v>', b'<f>A3+1</f><v>13</v>')
        with zipfile.ZipFile(path, 'w') as output:
            for name, data in entries.items():
                output.writestr(name, data)
        result = structured_excel.extract(path, 'inst1')
        table = result['structured_tables'][0]
        rows = {row['row']: {cell['coordinate']: cell for cell in row['cells']} for row in table['rows']}
        self.assertEqual(rows[2]['A2']['value'], 0)
        self.assertEqual(rows[2]['A2']['unit'], '万元')
        self.assertEqual(table['header_units'][0], '万元')
        self.assertIs(rows[2]['B2']['value'], False)
        self.assertEqual(rows[2]['C2']['value'], '00123')
        self.assertEqual(rows[2]['D2']['value'], .15)
        self.assertEqual(rows[2]['D2']['display'], '15%')
        self.assertEqual(rows[2]['D2']['unit'], '%')
        self.assertEqual(rows[2]['D2']['display_scale'], 100)
        self.assertEqual(rows[2]['E2']['formula'], '=A2+1')
        self.assertIsNone(rows[2]['E2']['value'])
        self.assertFalse(rows[2]['E2']['cached_available'])
        self.assertEqual(rows[3]['E3']['cached'], 13)
        self.assertTrue(rows[3]['E3']['cached_available'])
        self.assertTrue(rows[3]['F3']['inherited'])
        self.assertIsNone(rows[3]['F3']['value'])
        self.assertEqual(rows[3]['F3']['merged_anchor'], 'F2')
        self.assertTrue(table['rows'][2]['hidden'])
        self.assertEqual(rows[3]['A3']['unit'], '¥')
        self.assertEqual(table['row_count'], 3)
        self.assertEqual(table['header_columns'], [1, 2, 3, 4, 5, 6])

    def test_multiple_sheets_regions_multilevel_header_and_explicit_parallel_tables(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '统计'
        sheet['D1'] = '年度'
        sheet.merge_cells('D1:E1')
        sheet['D2'], sheet['E2'] = '类别', '金额'
        sheet['D3'], sheet['E3'] = 'A', 0
        sheet['D5'], sheet['E5'] = '类别2', '金额2'
        sheet['D6'], sheet['E6'] = 'B', 9
        second = book.create_sheet('并行')
        for row in [['名称', '金额', None, '编号', '数量'], ['A', 1, None, '001', 2], ['B', 3, None, '002', 4]]:
            second.append(row)
        second.add_table(Table(displayName='LeftTable', ref='A1:B3'))
        second.add_table(Table(displayName='RightTable', ref='D1:E3'))
        result = structured_excel.extract(self.save(book), 'inst1')
        tables = result['structured_tables']
        self.assertEqual(len(tables), 4)
        group = next(table for table in tables if table['sheet'] == '统计' and table['range'].startswith('D1'))
        self.assertEqual(group['header_rows'], [1, 2])
        self.assertEqual(group['header_columns'], [4, 5])
        self.assertEqual(group['headers'], ['年度 / 类别', '年度 / 金额'])
        self.assertTrue(group['rows'][1]['is_header'])
        self.assertFalse(group['rows'][2]['is_header'])
        self.assertEqual(sum(t['layout'] == 'defined-table' for t in tables), 2)
        right = next(t for t in tables if t['id'].endswith('RightTable'))
        self.assertEqual(right['header_columns'], [4, 5])
        self.assertEqual(right['rows'][1]['cells'][0]['coordinate'], 'D2')

    def test_stream_is_complete_and_completed_batches_resume_after_failure(self):
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet('数据')
        sheet.append(['编号', '金额'])
        for index in range(2300):
            sheet.append([f'{index:06d}', index])
        path = self.save(book)
        # Interrupt traversal after a committed 1000-row batch.
        with patch.object(structured_excel, 'MAX_CELLS', 2500):
            with self.assertRaisesRegex(ValueError, 'cell budget'):
                structured_excel.extract(path, 'inst1')
        result = structured_excel.extract(path, 'inst1')
        table = result['structured_tables'][0]
        self.assertEqual(table['mode'], 'artifact')
        self.assertEqual(len(table['rows']), 200)
        self.assertEqual(table['row_count'], 2301)
        self.assertGreaterEqual(table['reused_batches'], 1)
        artifact = artifact_store.resolve(table['artifact_id'], 'inst1')
        self.assertIsNotNone(artifact)
        with artifact.open() as handle:
            full = [json.loads(line) for line in handle]
        self.assertEqual(len(full), 2301)
        self.assertEqual(full[-1]['cells'][1]['value'], 2299)
        self.assertIsNone(artifact_store.resolve(table['artifact_id'], 'inst2'))
        replay = structured_excel.extract(path, 'inst1')['structured_tables'][0]
        self.assertEqual(replay['artifact_id'], table['artifact_id'])
        self.assertEqual(replay['reused_batches'], 3)

    def test_retry_only_selected_table_preserves_stable_ids_and_projection(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '原件'
        sheet.append(['类别', '金额'])
        sheet.append(['甲', 1])
        sheet.append([])
        sheet.append(['类别二', '金额二'])
        sheet.append(['乙', 2])
        path = self.save(book)
        original = structured_excel.extract(path, 'inst1')
        selected = original['structured_tables'][1]['id']
        retry = structured_excel.extract(path, 'inst1', [selected])
        self.assertEqual([t['id'] for t in retry['structured_tables']], [selected])
        self.assertEqual(retry['coverage'], {'total': 2, 'processed': 1, 'failed': 0, 'skipped': 1})
        self.assertEqual(retry['source_units'][1]['markdown'], retry['markdown'])
        self.assertEqual(retry['source_units'][1]['table_id'], selected)

    def test_xls_real_zero_boolean_merge_cached_only_capability(self):
        import xlwt
        book = xlwt.Workbook()
        sheet = book.add_sheet('旧表')
        for col, value in enumerate(['金额', '开关', '编号', '分类']):
            sheet.write(0, col, value)
        sheet.write(1, 0, 0)
        sheet.write(1, 1, False)
        sheet.write(1, 2, '00123')
        sheet.write(2, 0, 7)
        sheet.write_merge(1, 2, 3, 3, '父项')
        path = self.root / 'legacy.xls'
        book.save(str(path))
        table = structured_excel.extract(path, 'inst1')['structured_tables'][0]
        zero = table['rows'][1]['cells'][0]
        self.assertEqual(zero['value'], 0)
        self.assertEqual(zero['display'], '0')
        self.assertIs(table['rows'][1]['cells'][1]['value'], False)
        inherited = next(c for c in table['rows'][2]['cells'] if c['coordinate'] == 'D3')
        self.assertTrue(inherited['inherited'])
        self.assertIsNone(inherited['value'])
        self.assertEqual(table['formula_capability'], 'cached-only')

    def test_sparse_absolute_columns_explicit_summary_and_percent_text_scale(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['D1'], sheet['F1'] = '名称', '比例（%）'
        sheet['D2'], sheet['F2'] = 'A', .15
        sheet['F2'].number_format = '0%'
        sheet['D3'], sheet['F3'] = 'B', '20%'
        sheet['D4'], sheet['F4'] = '来源汇总', '=SUM(F2:F3)'
        table = structured_excel.extract(self.save(book), 'inst1')['structured_tables'][0]
        self.assertEqual(table['header_columns'], [4, 6])
        self.assertEqual(table['header_units'], [None, '%'])
        number = next(c for c in table['rows'][1]['cells'] if c['column'] == 6)
        text = next(c for c in table['rows'][2]['cells'] if c['column'] == 6)
        self.assertEqual(number['value'], .15)
        self.assertEqual(number['display_scale'], 100)
        self.assertEqual(text['value'], '20%')
        self.assertEqual(text['display_scale'], 1)
        self.assertTrue(table['rows'][-1]['is_summary'])

    def test_expanded_package_budget_checked_before_loading_workbook(self):
        path = self.root / 'bomb.xlsx'
        with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as output:
            output.writestr('xl/large.xml', 'x' * 2000)
        with patch.object(structured_excel, 'MAX_XML_BYTES', 1000), patch.object(openpyxl, 'load_workbook') as load:
            with self.assertRaisesRegex(ValueError, 'expanded input budget'):
                structured_excel.extract(path)
            load.assert_not_called()

    def test_merged_ranges_are_scoped_to_each_region_rectangle(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['A1'] = '标题'
        sheet.merge_cells('A1:C1')
        sheet['A3'], sheet['B3'], sheet['C3'] = 1, 2, 3
        sheet['A5'], sheet['B5'], sheet['C5'] = 4, 5, 6
        tables = structured_excel.extract(self.save(book), 'inst1')['structured_tables']
        with_merge = next(t for t in tables if t['range'] == 'A1:C1')
        without_merge = next(t for t in tables if t['range'] == 'A3:C3')
        self.assertEqual(with_merge['merged_ranges'], ['A1:C1'])
        # A region with no merges of its own must not claim its neighbour's span.
        self.assertEqual(without_merge['merged_ranges'], [])

    def test_single_merged_header_keeps_headers_unique_per_logical_column(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['A1'] = '数量'
        sheet.merge_cells('A1:C1')
        sheet['A2'], sheet['B2'], sheet['C2'] = 1, 2, 3
        table = structured_excel.extract(self.save(book), 'inst1')['structured_tables'][0]
        # One merged name, spanned columns without a name of their own fall back
        # to the column letter, so header lookup is unambiguous.
        self.assertEqual(table['headers'], ['数量', 'B', 'C'])
        self.assertEqual(len(set(table['headers'])), 3)

    def test_uncached_formula_is_not_indexed_as_document_content(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['A1'] = '列'
        sheet['A2'] = 1
        sheet['A3'] = '=SUM(A2:A2)'
        extracted = structured_excel.extract(self.save(book), 'inst1')
        table = extracted['structured_tables'][0]
        formula = next(c for c in table['rows'][-1]['cells'] if c['type'] == 'formula')
        self.assertIsNone(formula['value'])
        self.assertTrue(formula['cached_available'] is False)
        # Formula source stays traceable but must not become display content.
        self.assertEqual(formula['display'], '')
        self.assertEqual(formula['formula'], '=SUM(A2:A2)')
        self.assertTrue(table['warnings'])
        self.assertEqual(table['unresolved_formulas'], 1)
        self.assertNotIn('=SUM', extracted['markdown'])

    def test_xls_merged_cells_carry_range_and_formatted_display(self):
        import xlwt
        book = xlwt.Workbook()
        sheet = book.add_sheet('旧表')
        style = xlwt.XFStyle()
        sheet.write(1, 0, '数量')
        sheet.write_merge(0, 2, 2, 2, '父项')       # vertical merge C1:C3
        sheet.write(3, 0, 1234)
        sheet.write_merge(3, 3, 1, 3, 1234, style)  # horizontal merge B4:D4
        path = self.save_legacy(book, name='legacy-merge.xls')
        table = structured_excel.extract(path, 'inst1')['structured_tables'][0]
        anchor = next(c for c in table['rows'][3]['cells'] if c['coordinate'] == 'B4')
        inherited = next(c for c in table['rows'][3]['cells'] if c['coordinate'] == 'D4')
        self.assertEqual(anchor['merged_range'], 'B4:D4')
        self.assertEqual(anchor['merged_anchor'], 'B4')
        self.assertTrue(inherited['inherited'])
        self.assertEqual(inherited['merged_range'], 'B4:D4')
        self.assertEqual(inherited['merged_anchor'], 'B4')
        # The inherited display must be the anchor's formatted display, never a
        # raw str() of the cell value.
        self.assertEqual(inherited['display'], anchor['display'])
        self.assertNotIn('.0', inherited['display'])
        vertical = next(c for c in table['rows'][1]['cells'] if c['coordinate'] == 'C2')
        self.assertTrue(vertical['inherited'])
        self.assertEqual(vertical['display'], '父项')
        self.assertEqual(vertical['merged_range'], 'C1:C3')
        # Presentation only: an inherited cell never counts as a real fact, so
        # the merged value is reported exactly once.
        self.assertEqual(sum(1 for c in table['rows'][3]['cells'] if c['value'] is not None
                             and not c.get('inherited')), 2)

    def test_xls_merged_anchor_display_matches_inherited_display(self):
        import xlwt
        book = xlwt.Workbook()
        sheet = book.add_sheet('旧表')
        style = xlwt.XFStyle()
        style.num_format_str = 'General'
        sheet.write_merge(0, 0, 0, 2, 1234, style)
        table = structured_excel.extract(self.save_legacy(book, name='legacy-number.xls'), 'inst1')['structured_tables'][0]
        displays = [c['display'] for c in table['rows'][0]['cells']]
        self.assertEqual(displays, ['1234', '1234', '1234'])
