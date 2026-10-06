#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""引导 SOTA 主套件锚点语料库「系统测试-解析矩阵库」。

创建个人库并灌入 ≥10 篇锚点文档（覆盖 P2-01..04 的锚点事实、P2-05 的八章考勤手册、
P3-01 夏令时、P3-02 冲突语料、C 系列探针的第十条/N+1 等），并上传多 Sheet xlsx。
等待全部发布后退出。用法：python3 bootstrap_anchor_corpus.py
"""
import io
import json
import os
import ssl
import sys
import time
import urllib.error
import urllib.request

API = os.environ.get("API_BASE", "http://127.0.0.1:3202").rstrip("/")
USER = os.environ.get("LLMWIKI_USER", "admin")
PASS = os.environ.get("LLMWIKI_PASS", "admin123")
KB_NAME = "系统测试-解析矩阵库"
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE


def http(method, path, body=None, token=None, raw_body=None, headers=None, timeout=120):
    url = f"{API}{path}"
    data = raw_body if raw_body is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(url, data=data, method=method)
    if raw_body is None:
        req.add_header("Content-Type", "application/json")
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=timeout, context=CTX) as r:
            return r.status, r.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8", "replace")


def login():
    st, raw = http("POST", "/api/v1/auth/login", {"username": USER, "password": PASS})
    assert st == 200, f"login failed {st}: {raw[:200]}"
    payload = json.loads(raw)
    return payload.get("token") or payload["accessToken"]


def find_kb(token):
    page = 1
    while page <= 10:
        st, raw = http("GET", f"/api/v1/kbs?page={page}&limit=100", token=token)
        if st != 200:
            return None
        payload = json.loads(raw)
        items = payload.get("items") if isinstance(payload, dict) else payload
        for kb in items or []:
            if kb.get("name") == KB_NAME:
                return kb["id"]
        total = int(payload.get("total") or 0) if isinstance(payload, dict) else 0
        if page * 100 >= total or not items:
            return None
        page += 1
    return None


# ---------------------------------------------------------------- 锚点文档
def attendance_detailed():
    return """# 企业考勤管理制度详细手册

## 第一章 总则
为规范公司考勤管理，维护工作秩序，提升工作效率，依据国家有关法律法规，结合公司实际情况，制定本手册。本手册适用于公司全体正式员工与实习生。

## 第二章 工时制度与作息时间
公司实行标准工时制。常规作息时间为：工作日上午 09:00 上班，18:00 下班。
每年夏季（6 月 1 日至 9 月 30 日）实行夏令时作息：上午上班时间调整为 **08:30**，下班时间 17:30，午休相应提前。

## 第三章 考勤方式与规范
员工上下班须通过企业办公系统打卡签到签退。每人每日打卡两次，漏打卡须于当日在系统内提交补卡申请。

## 第四章 考勤异常认定与处理
迟到：超过规定上班时间打卡。早退：早于规定下班时间打卡。旷工：未经请假擅自缺勤或当日无打卡记录且无补卡申请。

## 第五章 加班管理
加班须事先审批。工作日加班按 1.5 倍计算，休息日按 2 倍，法定节假日按 3 倍。

## 第六章 请假休假管理
请假类型包括事假、病假、年休假、婚假、产假、陪产假、丧假等。年休假按司龄分级：满 1 年 5 天，满 10 年 10 天。

## 第七章 考勤统计与薪资核算
人力资源部每月 3 个工作日内完成上月考勤统计，考勤异常结果计入当月绩效。

## 第八章 附则
本手册由人力资源部负责解释，自发布之日起施行。
"""


def attendance_old():
    return """# 考勤管理规定（2024 旧版）

第一条 公司实行每日 8 小时工作制，上班时间为上午 **09:00**，下班时间为 18:00。
第二条 夏季（7-8 月）作息：上班时间 09:00 不变，午休延长至 14:00。
第三条 员工忘打卡每月补卡不得超过 2 次。
第四条 连续旷工 3 天或年度累计旷工 7 天的，视为严重违反规章制度。
"""


def inspection_table():
    rows = ["| 设备编号 | 设备名称 | 巡检周期（天） | 责任部门 | 考核分值 |",
            "| --- | --- | --- | --- | --- |"]
    anchors = {"EQ-0077": ("空压机组", 30), "EQ-0102": ("中央空调机组", 15),
               "EQ-0233": ("高压配电柜", 7), "EQ-0310": ("电梯系统", 14)}
    for i in range(1, 81):
        eq = f"EQ-{i:04d}"
        name, cycle = anchors.get(eq, (f"测试设备{i}号", 30 + i))
        rows.append(f"| {eq} | {name} | {cycle} | 设备部 | {90 - i} |")
    return "# 设备巡检考核表（big_table）\n\n本表规定各设备的巡检周期与考核分值。\n\n" + "\n".join(rows) + "\n"


def big_doc():
    parts = ["# 平台运维知识大全文档"]
    for i in range(1, 320):
        parts.append(f"\n## 运维条目 {i:03d}\n常见故障场景 {i}：系统日志出现异常波动时，应首先检查第 {i} 号监控指标与最近一次变更记录，确认是否存在资源瓶颈或配置漂移，并按照标准应急预案执行回滚或扩容操作。")
    parts.append("\n## 文档校验锚点\n本大文档的校验锚点编号为 **BIGDOC-VERIFY=7788**，用于验证大文档尾部检索能力。\n")
    return "\n".join(parts)


def ultralong_budget():
    parts = ["# 天穹-2026 年度项目总体规划书",
             "\n## 项目总预算\n天穹-2026 项目的总预算为 **3.75 亿元**，分三个阶段拨付。"]
    for i in range(1, 260):
        parts.append(f"\n### 预算科目 {i:03d}\n本科目涵盖阶段{(i % 3) + 1}相关的人力投入、外购服务与设备采购预算，金额按季度滚动核对，偏差超过 5% 时触发预算评审流程。")
    return "\n".join(parts)


def employee_handbook_rules():
    return """# 员工手册（纪律与补偿条款）

第十条 经济补偿：公司依法解除劳动合同时，按员工在本单位工作年限支付经济补偿，每满一年支付一个月工资，即 **N+1** 标准执行（N 为工作年限，+1 为代通知金）。
第十一条 保密义务：员工在职期间及离职后两年内不得泄露公司商业秘密。
第十二条 廉洁纪律：员工不得收受供应商财物，违者一经查实立即解除劳动合同。
"""


def misc_docs():
    return [
        ("信息安全管理制度", "# 信息安全管理制度\n\n第一条 公司信息系统实行分级保护，核心数据加密存储。\n第二条 员工账号实行最小权限原则，离职当日回收全部权限。\n第三条 每季度开展一次安全审计。"),
        ("差旅费报销规定", "# 差旅费报销规定\n\n一、市内交通费据实报销，单次上限 200 元。\n二、住宿标准：一线城市 600 元/晚，其他城市 450 元/晚。\n三、报销须于行程结束后 10 个工作日内提交。"),
        ("会议管理制度", "# 会议管理制度\n\n一、全员会议须提前 24 小时发出议程。\n二、会议纪要须于会后 1 个工作日内归档至知识库。\n三、迟到 10 分钟以上计入考勤异常。"),
        ("采购流程规范", "# 采购流程规范\n\n1. 单笔 1 万元以下由部门负责人审批。\n2. 单笔 1 万至 10 万元须分管副总审批。\n3. 单笔 10 万元以上须招标小组评议。"),
    ]


def build_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "考核汇总"
    ws1.append(["考核汇总表", ""])
    ws1.append(["汇总编号", "SUM-2026-5566"])
    ws1.append(["考核周期", "2026 年度 Q1-Q4"])
    ws1.append(["参与部门数", 12])
    ws1.append(["平均分", 91.4])
    ws2 = wb.create_sheet("部门明细")
    ws2.append(["部门", "季度", "得分"])
    for d, s in [("研发中心", 93.2), ("市场部", 90.1), ("合规部", 94.7), ("设备部", 88.5)]:
        ws2.append([d, "Q1", s])
    ws3 = wb.create_sheet("指标说明")
    ws3.append(["指标", "权重", "说明"])
    ws3.append(["合规性", 0.4, "制度执行与审计扣分"])
    ws3.append(["产出质量", 0.4, "交付物抽检合格率"])
    ws3.append(["协作", 0.2, "跨部门评价"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    token = login()
    print("[1] logged in")
    kb_id = find_kb(token)
    if not kb_id:
        st, raw = http("POST", "/api/v1/kbs/personal",
                       {"name": KB_NAME, "description": "SOTA 主套件锚点语料（解析矩阵：表格/大文档/超长文档/多Sheet xlsx）"}, token=token)
        assert st in (200, 201), f"create kb failed {st}: {raw[:300]}"
        kb_id = json.loads(raw)["id"] if "id" in json.loads(raw) else json.loads(raw).get("knowledgeBase", {}).get("id")
        print(f"[2] KB created {kb_id}")
    else:
        print(f"[2] KB exists {kb_id}")

    docs = [
        ("企业考勤管理制度详细手册", attendance_detailed()),
        ("考勤管理规定（2024 旧版）", attendance_old()),
        ("设备巡检考核表（big_table）", inspection_table()),
        ("平台运维知识大全文档", big_doc()),
        ("天穹-2026 年度项目总体规划书", ultralong_budget()),
        ("员工手册（纪律与补偿条款）", employee_handbook_rules()),
    ] + misc_docs()

    st, raw = http("GET", f"/api/v1/kbs/{kb_id}/documents", token=token)
    existing = set()
    if st == 200:
        payload = json.loads(raw)
        items = payload if isinstance(payload, list) else payload.get("items") or payload.get("documents") or []
        existing = {d.get("title") for d in items}

    for title, content in docs:
        if title in existing:
            print(f"    skip(existing) {title}")
            continue
        st, raw = http("POST", f"/api/v1/kbs/{kb_id}/documents/text",
                       {"title": title, "content": content, "duplicateMode": "skip"}, token=token, timeout=60)
        print(f"    add [{st}] {title}" + ("" if st in (200, 201) else f" {raw[:120]}"))

    # 多 Sheet xlsx 上传（ASCII 文件名：服务端对非 ASCII Content-Disposition 兼容有限）
    if "exam-summary-2026.xlsx" not in existing:
        xlsx = build_xlsx()
        boundary = "----gbkbootstrap1234"
        fname = "exam-summary-2026.xlsx"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{fname}\"\r\n"
                f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() \
               + xlsx + f"\r\n--{boundary}--\r\n".encode()
        st, raw = http("POST", f"/api/v1/kbs/{kb_id}/documents", raw_body=body, timeout=60,
                       headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        print(f"    xlsx [{st}] " + ("" if st in (200, 201) else raw[:200]))

    # 等待发布
    print("[3] waiting for publish ...")
    deadline = time.time() + 300
    while time.time() < deadline:
        st, raw = http("GET", f"/api/v1/kbs/{kb_id}/documents", token=token)
        payload = json.loads(raw) if st == 200 else {}
        items = payload if isinstance(payload, list) else payload.get("items") or payload.get("documents") or []
        published = sum(1 for d in items if d.get("status") == "published")
        total = len(items)
        print(f"    published {published}/{total}")
        if total >= 11 and published == total:
            print("[OK] corpus ready")
            return 0
        time.sleep(10)
    print("[WARN] timeout waiting publish")
    return 1


if __name__ == "__main__":
    sys.exit(main())
